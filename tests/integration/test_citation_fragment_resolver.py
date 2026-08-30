from __future__ import annotations

import base64
import hashlib
import io
import os
import tempfile
import unittest

from docx import Document

from lvke_mcp.adapters.source_files_repository import (
    SourceFileError,
    resolve_citation_fragment,
)
from lvke_mcp.servers.lvke_data_acquisition import service as acquisition
from lvke_mcp.servers.lvke_source_files import service as source_files


def _minimal_text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> /MediaBox [0 0 200 200] /Contents 5 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(output)
    output += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    for offset in offsets:
        output += f"{offset:010d} 00000 n \n".encode()
    output += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref_at}\n%%EOF\n".encode()
    return bytes(output)


class CitationFragmentResolverTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tempdir = tempfile.TemporaryDirectory(prefix="lvke-citation-resolver-")
        self.previous = os.environ.get("LVKE_MCP_DATA_DIR")
        os.environ["LVKE_MCP_DATA_DIR"] = self.tempdir.name
        self.workspace = "citation-resolver-test"

    def tearDown(self) -> None:
        if self.previous is None:
            os.environ.pop("LVKE_MCP_DATA_DIR", None)
        else:
            os.environ["LVKE_MCP_DATA_DIR"] = self.previous
        self.tempdir.cleanup()

    def _import(self, name: str, mime: str, content: bytes) -> tuple[str, str]:
        result = source_files.import_content(
            workspace_id=self.workspace,
            original_filename=name,
            declared_mime=mime,
            content_base64=base64.b64encode(content).decode("ascii"),
            idempotency_key=f"citation-{name}",
            parse_immediately=True,
        )
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["source_file"]["status"], "succeeded", result)
        return result["file_id"], "sha256:" + hashlib.sha256(content).hexdigest()

    def test_resolves_pdf_page_and_offsets(self) -> None:
        file_id, source_hash = self._import(
            "source.pdf", "application/pdf", _minimal_text_pdf("Revenue 123 units")
        )
        result = resolve_citation_fragment(
            self.workspace,
            source_id=file_id,
            source_hash=source_hash,
            locator={"kind": "pdf_page", "page": 1, "start_offset": 8, "end_offset": 11},
        )
        self.assertEqual(result["fragment_text"], "123")
        self.assertEqual(result["locator"]["locator"], "pdf_page:1")
        self.assertEqual(result["fragment_hash"], "sha256:" + hashlib.sha256(b"123").hexdigest())

    def test_resolves_csv_cell_by_a1_and_row_column(self) -> None:
        file_id, source_hash = self._import(
            "source.csv", "text/csv", b"field,value\nrevenue,5200\n"
        )
        by_a1 = resolve_citation_fragment(
            self.workspace,
            source_id=file_id,
            source_hash=source_hash,
            locator={"kind": "csv_cell", "cell": "B2"},
        )
        by_position = resolve_citation_fragment(
            self.workspace,
            source_id=file_id,
            source_hash=source_hash,
            locator="csv:2:2",
        )
        self.assertEqual(by_a1["fragment_text"], "5200")
        self.assertEqual(by_a1["fragment_hash"], by_position["fragment_hash"])
        self.assertEqual(by_position["locator"]["cell"], "B2")

    def test_resolves_docx_paragraph_and_text_document(self) -> None:
        document = Document()
        document.add_paragraph("First governed paragraph")
        buffer = io.BytesIO()
        document.save(buffer)
        docx_id, docx_hash = self._import(
            "source.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            buffer.getvalue(),
        )
        paragraph = resolve_citation_fragment(
            self.workspace,
            source_id=docx_id,
            source_hash=docx_hash,
            locator="paragraph:1",
        )
        self.assertEqual(paragraph["fragment_text"], "First governed paragraph")

        text_id, text_hash = self._import("source.txt", "text/plain", b"alpha beta gamma")
        fragment = resolve_citation_fragment(
            self.workspace,
            source_id=text_id,
            source_hash=text_hash,
            locator={"kind": "document_text", "start_offset": 6, "end_offset": 10},
            supplied_fragment="beta",
        )
        self.assertEqual(fragment["fragment_text"], "beta")

    def test_resolves_stored_spreadsheet_cell_locator(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inputs"
        sheet["B2"] = 5200
        buffer = io.BytesIO()
        workbook.save(buffer)
        file_id, source_hash = self._import(
            "source.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            buffer.getvalue(),
        )
        result = resolve_citation_fragment(
            self.workspace,
            source_id=file_id,
            source_hash=source_hash,
            locator="workbook:Inputs!B2",
        )
        self.assertEqual(result["fragment_text"], "5200")
        self.assertEqual(result["locator"]["locator"], "workbook:Inputs!B2")

    def test_resolves_snapshot_body_not_artifact_record_hash(self) -> None:
        content = "政策正文含可核对的建设时序。"
        imported = acquisition.import_external_snapshot(
            self.workspace,
            url="https://example.com/policy",
            title="policy",
            content=content,
            provider="codex-browser",
            provider_tool="browser_snapshot",
            retrieved_at="2026-08-29T12:00:00+08:00",
            content_kind="extracted_full_text",
        )
        result = resolve_citation_fragment(
            self.workspace,
            source_id=imported["source_snapshot_id"],
            source_hash=imported["external_content_hash"],
            locator={"kind": "web_snapshot", "start_offset": 5, "end_offset": 9},
        )
        self.assertEqual(result["fragment_text"], "可核对的")
        self.assertNotEqual(imported["content_hash"], result["source_hash"])

    def test_rejects_out_of_bounds_hash_fragment_and_cross_workspace(self) -> None:
        file_id, source_hash = self._import("guard.txt", "text/plain", b"immutable text")
        cases = [
            ({"locator": {"kind": "document_text", "start_offset": 0, "end_offset": 99}}, "citation_locator_out_of_bounds"),
            ({"source_hash": "sha256:" + "0" * 64}, "citation_source_hash_mismatch"),
            ({"supplied_fragment": "tampered"}, "citation_fragment_mismatch"),
            ({"supplied_fragment_hash": "sha256:" + "f" * 64}, "citation_fragment_hash_mismatch"),
        ]
        for overrides, code in cases:
            with self.subTest(code=code), self.assertRaises(SourceFileError) as captured:
                resolve_citation_fragment(
                    self.workspace,
                    source_id=file_id,
                    source_hash=overrides.get("source_hash", source_hash),
                    locator=overrides.get("locator", "document_text"),
                    supplied_fragment=overrides.get("supplied_fragment", ""),
                    supplied_fragment_hash=overrides.get("supplied_fragment_hash", ""),
                )
            self.assertEqual(captured.exception.detail["code"], code)
        with self.assertRaises(SourceFileError) as captured:
            resolve_citation_fragment(
                "another-workspace",
                source_id=file_id,
                source_hash=source_hash,
                locator="document_text",
            )
        self.assertEqual(captured.exception.detail["code"], "citation_source_not_found")


if __name__ == "__main__":
    unittest.main()
