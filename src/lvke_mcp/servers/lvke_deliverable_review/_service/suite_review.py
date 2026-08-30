"""Seven-dimension research-suite review objects and lifecycle operations."""

from __future__ import annotations

import hashlib
import csv
import json
import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from lvke_mcp.runtime.storage import require_safe_id, sha256_json, utc_now
from lvke_mcp.runtime.formal_promotion import FormalLineageError
from lvke_mcp.servers.lvke_deliverable_review import rules
from lvke_mcp.servers.lvke_deliverable_review.contracts import (
    COMPLIANCE_STATUSES,
    DIMENSION_STATUSES,
    FULL_SUITE_REQUIRED_ROLES,
    REVIEW_COMPONENT_ROLES,
    REVIEW_DIMENSIONS,
    REVIEW_MODES,
    REVIEW_PROFILES,
    SEVERITY_ORDER,
    finding_blocks,
)
from lvke_mcp.servers.lvke_deliverable_review.store import STORE

from .base import (
    DIMENSION_CONFIRMATION_STORE,
    DIMENSION_RESULT_STORE,
    DOSSIER_STORE,
    EXTRACTION_CONFIRMATION_STORE,
    PACKAGE_DRAFT_STORE,
    REVIEW_PACKAGE_STORE,
    SUITE_ASSESSMENT_STORE,
    _blocked,
    _message,
    _ok,
    _write,
)
from .suite_package import (
    get_package,
    internal_component as _build_internal_component,
    internal_package_lineage as _validate_internal_package_lineage,
    package_integrity_reasons as _validate_package_integrity,
    source_records as _source_records,
    verified_source_component as _verified_source_component,
)


CHECK_CATALOG: dict[str, dict[str, str]] = {
    "COMP.REQUIREMENT.COVERAGE": {"dimension": "compliance", "kind": "deterministic"},
    "COMP.SUBSTANTIVE.REVIEW": {"dimension": "compliance", "kind": "semantic"},
    "ARTICLE.STRUCTURE.PLACEHOLDER": {"dimension": "article_quality", "kind": "deterministic"},
    "ARTICLE.DUPLICATE.TEMPLATE": {"dimension": "article_quality", "kind": "deterministic"},
    "ARTICLE.LANGUAGE.LOGIC": {"dimension": "article_quality", "kind": "semantic"},
    "ARTICLE.VISUAL.LAYOUT": {"dimension": "article_quality", "kind": "semantic"},
    "DATA.SCHEMA.PROFILE": {"dimension": "data_quality", "kind": "deterministic"},
    "DATA.ANOMALY.RECONCILIATION": {"dimension": "data_quality", "kind": "semantic"},
    "SOURCE.IDENTITY.LOCATOR": {"dimension": "source_quality", "kind": "deterministic"},
    "SOURCE.CLAIM.SUPPORT": {"dimension": "source_quality", "kind": "semantic"},
    "FINMODEL.WORKBOOK.INTEGRITY": {"dimension": "financial_model", "kind": "deterministic"},
    "FINMODEL.ASSUMPTION.LOGIC": {"dimension": "financial_model", "kind": "semantic"},
    "FINTABLE.RECONCILIATION": {"dimension": "financial_tables", "kind": "deterministic"},
    "FINTABLE.SUBSTANTIVE.REVIEW": {"dimension": "financial_tables", "kind": "semantic"},
    "FEASIBILITY.STRUCTURE.COVERAGE": {"dimension": "feasibility", "kind": "deterministic"},
    "FEASIBILITY.DECISION.LOGIC": {"dimension": "feasibility", "kind": "semantic"},
}


_ROLE_TERMS = {
    "finance_tables": ("十三表", "十五表", "财务表", "附表", "tables", "statement"),
    "finance_model": ("财务模型", "测算模型", "model", "测算表", "估值模型"),
    "base_data": ("基础数据", "原始数据", "数据表", "dataset", "台账", "清单"),
    "source_evidence": ("来源", "证据", "政策", "法规", "批复", "合同", "source", "evidence"),
    "report": ("可研", "研究报告", "报告", "研报", "feasibility", "report"),
}


def _hash(value: str) -> str:
    text = str(value or "").lower().strip()
    return text if text.startswith("sha256:") else f"sha256:{text}" if text else ""


def _suggest_role(component: dict[str, Any]) -> tuple[str, float, list[str]]:
    filename = str(component.get("filename") or "").lower()
    mime = str(component.get("mime_type") or "").lower()
    suffix = Path(filename).suffix.lower()
    candidates: list[tuple[float, str, str]] = []
    for role, terms in _ROLE_TERMS.items():
        for term in terms:
            if term.lower() in filename:
                candidates.append((0.92, role, f"filename:{term}"))
    if suffix in {".xlsx", ".xls", ".xlsm"} or "spreadsheet" in mime or "excel" in mime:
        candidates.append((0.70, "finance_model", "spreadsheet_format"))
    if suffix in {".csv", ".tsv", ".json", ".jsonl"}:
        candidates.append((0.75, "base_data", "tabular_format"))
    if suffix in {".docx", ".pdf", ".md", ".txt"}:
        candidates.append((0.60, "report", "document_format"))
    if not candidates:
        return "attachment", 0.40, ["unclassified"]
    candidates.sort(reverse=True)
    confidence, role, reason = candidates[0]
    return role, confidence, [reason]


def _internal_component(
    workspace_id: str,
    target: dict[str, Any],
) -> tuple[dict[str, Any] | None, list[str]]:
    from .target_resolve import _resolve_target
    return _build_internal_component(
        workspace_id,
        target,
        resolve_target=_resolve_target,
    )


def _internal_package_lineage(
    workspace_id: str,
    components: list[dict[str, Any]],
) -> dict[str, Any]:
    from .target_resolve import _resolve_target

    return _validate_internal_package_lineage(
        workspace_id,
        components,
        resolve_target=_resolve_target,
    )


def _component_roles(components: list[dict[str, Any]]) -> set[str]:
    return {
        str(component.get("role") or component.get("suggested_role") or "")
        for component in components
    }


def prepare_package(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        review_mode = str(args.get("review_mode") or "external")
        profile = str(args.get("review_profile") or "standard")
        if review_mode not in REVIEW_MODES:
            return _blocked("review_mode_invalid", "review_mode 必须为 internal 或 external")
        if profile not in REVIEW_PROFILES:
            return _blocked("review_profile_invalid", "review_profile 必须为 quick、standard 或 deep")
        components: list[dict[str, Any]] = []
        blockers: list[str] = []
        for file_id in args.get("source_file_ids") or []:
            component, reasons = _verified_source_component(workspace_id, str(file_id))
            if reasons or component is None:
                blockers.extend(f"{file_id}:{reason}" for reason in reasons)
                continue
            role, confidence, reasons = _suggest_role(component)
            components.append({
                **component,
                "suggested_role": role,
                "classification_confidence": confidence,
                "classification_reasons": reasons,
            })
        for target in args.get("internal_targets") or []:
            component, reasons = _internal_component(workspace_id, target)
            if reasons or component is None:
                target_id = str((target or {}).get("target_id") or "unknown")
                blockers.extend(f"{target_id}:{reason}" for reason in reasons)
                continue
            components.append(component)
        if blockers:
            return _blocked(blockers[0], "套件组件无法完整解析", blockers=sorted(set(blockers)))
        if not components:
            return _blocked("review_package_components_required", "至少提供一个 SourceFile 或内部不可变对象")
        ids = [str(row.get("component_id") or "") for row in components]
        if len(ids) != len(set(ids)):
            return _blocked("review_package_component_duplicate", "套件组件不得重复")
        roles = _component_roles(components)
        missing = sorted(FULL_SUITE_REQUIRED_ROLES - roles)
        extraction_pending = sorted(
            str(row.get("component_id") or "")
            for row in components
            if str(row.get("ocr_status") or "") in {"needs_ocr", "pending", "failed"}
            or str((row.get("analysis_summary") or {}).get("degraded_reason") or "")
        )
        payload = {
            "review_mode": review_mode,
            "review_profile": profile,
            "project_scope": deepcopy(args.get("project_scope") or {}),
            "components": sorted(components, key=lambda row: str(row.get("component_id") or "")),
            "missing_required_roles": missing,
            "full_suite_candidate": not missing,
            "extraction_confirmation_required": extraction_pending,
            "classification_method": "review-package-classifier.v1",
        }
        record = PACKAGE_DRAFT_STORE.put(
            workspace_id,
            payload,
            producer="lvke-deliverable-review.review_package_prepare",
            source_ids=ids,
            basis=payload,
            schema_version="ReviewPackageDraft.v1",
        )
        return _ok(
            review_package_draft_id=record["object_id"],
            draft_hash=record["content_hash"],
            components=payload["components"],
            missing_required_roles=missing,
            full_suite_candidate=not missing,
            extraction_confirmation_required=extraction_pending,
            resource_uris=[record["resource_uri"]],
            warnings=["文件角色必须经 review_package_confirm 显式确认"],
            blockers=[],
            next_actions=["确认每个组件角色后调用 review_package_confirm"],
        )
    return _write("review_package_prepare", args, execute)


def _verified_record(store: Any, workspace_id: str, object_id: str) -> dict[str, Any] | None:
    try:
        record = store.get(workspace_id, require_safe_id(object_id, "object_id"))
    except (OSError, ValueError):
        return None
    if not isinstance(record, dict) or not isinstance(record.get("payload"), dict):
        return None
    if record.get("content_hash") != sha256_json(record["payload"]):
        return None
    if record.get("basis_hash") != sha256_json(record.get("basis")):
        return None
    return record


def confirm_package(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        draft_id = str(args.get("review_package_draft_id") or "")
        draft = _verified_record(PACKAGE_DRAFT_STORE, workspace_id, draft_id)
        if draft is None:
            return _blocked("review_package_draft_not_found", "套件草稿不存在或完整性无效")
        if str(args.get("expected_draft_hash") or "") != str(draft.get("content_hash") or ""):
            return _blocked("review_package_draft_hash_mismatch", "套件草稿 hash 与确认请求不一致")
        payload = draft["payload"]
        role_map = {
            str(row.get("component_id") or ""): str(row.get("role") or "")
            for row in args.get("component_roles") or []
        }
        expected_ids = {str(row.get("component_id") or "") for row in payload.get("components") or []}
        if set(role_map) != expected_ids:
            return _blocked("review_package_role_mapping_incomplete", "必须为草稿中的每个组件确认且只确认一个角色")
        if any(role not in REVIEW_COMPONENT_ROLES for role in role_map.values()):
            return _blocked("review_package_role_invalid", "套件组件角色不受支持")
        components: list[dict[str, Any]] = []
        blockers: list[str] = []
        for raw in payload.get("components") or []:
            component = deepcopy(raw)
            if component.get("component_type") == "source_file":
                current, reasons = _verified_source_component(
                    workspace_id, str(component.get("component_id") or ""),
                )
                if reasons or current is None:
                    blockers.extend(reasons or ["source_component_unavailable"])
                    continue
                for field in ("content_hash", "analysis_hash", "parse_status", "ocr_status"):
                    if str(current.get(field) or "") != str(component.get(field) or ""):
                        blockers.append(f"source_component_changed:{component.get('component_id')}:{field}")
            component["role"] = role_map[str(component.get("component_id") or "")]
            components.append(component)
        if blockers:
            return _blocked(blockers[0], "确认前套件组件已变化", blockers=sorted(set(blockers)))
        roles = _component_roles(components)
        missing = sorted(FULL_SUITE_REQUIRED_ROLES - roles)
        canonical_lineage: dict[str, Any] = {}
        if payload["review_mode"] == "internal":
            try:
                canonical_lineage = _internal_package_lineage(workspace_id, components)
            except FormalLineageError as exc:
                return _blocked(exc.code, f"内部套件正式谱系校验失败：{exc.message}")
        package_payload = {
            "draft_id": draft_id,
            "draft_hash": draft["content_hash"],
            "review_mode": payload["review_mode"],
            "review_profile": payload["review_profile"],
            "project_scope": deepcopy(payload.get("project_scope") or {}),
            "components": sorted(components, key=lambda row: str(row.get("component_id") or "")),
            "present_roles": sorted(roles),
            "missing_required_roles": missing,
            "full_suite": not missing,
            "extraction_confirmation_required": list(payload.get("extraction_confirmation_required") or []),
            "role_confirmation": {
                "confirmed": True,
                "statement": str(args.get("confirmation_statement") or ""),
            },
            **canonical_lineage,
        }
        record = REVIEW_PACKAGE_STORE.put(
            workspace_id,
            package_payload,
            producer="lvke-deliverable-review.review_package_confirm",
            source_ids=[draft_id, *sorted(expected_ids)],
            basis={"draft_hash": draft["content_hash"], "component_roles": role_map},
            schema_version="ReviewPackage.v1",
        )
        return _ok(
            review_package_id=record["object_id"],
            review_package_hash=record["content_hash"],
            full_suite=not missing,
            missing_required_roles=missing,
            components=package_payload["components"],
            resource_uris=[record["resource_uri"]],
            warnings=[] if not missing else ["材料不满足完整研报套件定义，只能形成专项结论"],
            blockers=[],
            next_actions=["需要 OCR/低置信度确认时调用 review_confirm_extraction，然后调用 review_prepare"],
        )
    return _write("review_package_confirm", args, execute)


def confirm_extraction(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        package_id = str(args.get("review_package_id") or "")
        package = _verified_record(REVIEW_PACKAGE_STORE, workspace_id, package_id)
        if package is None:
            return _blocked("review_package_not_found", "审查套件不存在或完整性无效")
        integrity_reasons = package_integrity_reasons(workspace_id, package)
        if integrity_reasons:
            return _blocked(
                integrity_reasons[0],
                "提取确认前 ReviewPackage 完整性或正式谱系校验失败",
                blockers=integrity_reasons,
            )
        component_ids = {
            str(row.get("component_id") or "")
            for row in package["payload"].get("components") or []
        }
        from lvke_mcp.adapters.source_files_repository import SourceFileError, resolve_citation_fragment

        confirmations: list[dict[str, Any]] = []
        for row in args.get("confirmations") or []:
            source_id = str(row.get("source_id") or "")
            if source_id not in component_ids:
                return _blocked("extraction_source_not_in_package", "OCR 确认来源不属于当前套件")
            try:
                resolved = resolve_citation_fragment(
                    workspace_id,
                    source_id=source_id,
                    locator=row.get("locator"),
                    source_hash=row.get("source_hash"),
                    supplied_fragment=row.get("fragment_text", ""),
                    supplied_fragment_hash=row.get("fragment_hash", ""),
                )
            except SourceFileError as exc:
                return _blocked(str(exc.detail.get("code") or "extraction_fragment_invalid"), str(exc.detail.get("message") or "片段验证失败"))
            confirmations.append({
                **resolved,
                "source_id": source_id,
                "confirmation_kind": str(row.get("confirmation_kind") or "ocr_fragment"),
                "confirmed_value": deepcopy(row.get("confirmed_value")),
                "note": str(row.get("note") or ""),
            })
        if not confirmations:
            return _blocked("extraction_confirmations_required", "至少确认一个可验证片段")
        confirmation_payload = {
            "review_package_id": package_id,
            "review_package_hash": package["content_hash"],
            "confirmations": confirmations,
        }
        record = EXTRACTION_CONFIRMATION_STORE.put(
            workspace_id,
            confirmation_payload,
            producer="lvke-deliverable-review.review_confirm_extraction",
            source_ids=[package_id, *sorted({row["source_id"] for row in confirmations})],
            basis=confirmation_payload,
            schema_version="ReviewExtractionConfirmation.v1",
        )
        return _ok(
            extraction_confirmation_id=record["object_id"],
            confirmed_count=len(confirmations),
            resource_uris=[record["resource_uri"]],
            warnings=["该记录仅确认提取片段，不证明片段语义支持任何结论"],
            blockers=[],
            next_actions=["调用 review_prepare 审查该 ReviewPackage"],
        )
    return _write("review_confirm_extraction", args, execute)


def package_integrity_reasons(workspace_id: str, package_record: dict[str, Any]) -> list[str]:
    from .target_resolve import _resolve_target

    return _validate_package_integrity(
        workspace_id,
        package_record,
        resolve_target=_resolve_target,
    )


def _internal_component_text(workspace_id: str, component: dict[str, Any]) -> str:
    """Read reviewable text from an internal immutable report component.

    此前只从 ``source_file`` 组件取正文，于是内部对象构成的套件（零材料预览、
    可研链自有 report_revision）永远命中 ``report_text_unavailable``，
    ARTICLE.* 与 FEASIBILITY.STRUCTURE.COVERAGE 于是恒判不通过——报告明明在套件
    里，审查却说"拿不到正文"。那是读取路径缺一支，不是交付缺陷。
    """

    component_type = str(component.get("component_type") or "")
    component_id = str(component.get("component_id") or "")
    if not component_id:
        return ""
    if component_type == "report_artifact":
        try:
            from lvke_mcp.adapters.zero_material_repository import REPORT_STORE

            record = REPORT_STORE.get(workspace_id, component_id)
        except (OSError, ValueError):
            return ""
        payload = record.get("payload") if isinstance(record, dict) else None
        if isinstance(payload, dict):
            return str(payload.get("content_markdown") or "")
        return ""
    if component_type == "report_revision":
        try:
            from lvke_mcp.adapters.report_repository import REVISION_STORE

            record = REVISION_STORE.get(workspace_id, component_id)
        except (OSError, ValueError):
            return ""
        payload = record.get("payload") if isinstance(record, dict) else None
        if isinstance(payload, dict):
            return str(payload.get("content") or payload.get("content_markdown") or "")
    return ""


def _package_texts(workspace_id: str, package: dict[str, Any], role: str) -> list[tuple[str, str]]:
    _files, analyses = _source_records(workspace_id)
    rows: list[tuple[str, str]] = []
    for component in package.get("components") or []:
        if component.get("role") != role:
            continue
        component_id = str(component.get("component_id") or "")
        if component.get("component_type") == "source_file":
            analysis = analyses.get(component_id) if isinstance(analyses.get(component_id), dict) else {}
            text = str((analysis or {}).get("text_preview") or "")
        else:
            text = _internal_component_text(workspace_id, component)
        if text:
            rows.append((component_id, text))
    return rows


def _suite_finding(
    check_id: str,
    severity: str,
    message: str,
    *,
    location: dict[str, Any],
    actual: Any = None,
    remediation: str = "补充或修订材料后创建新套件并复测",
) -> dict[str, Any]:
    dimension = CHECK_CATALOG[check_id]["dimension"]
    return rules.finding(
        check_id,
        severity,
        message,
        category=dimension,
        target_location=location,
        actual=actual,
        review_area=dimension,
        remediation=remediation,
    )


def _csv_profile(path: Path) -> tuple[list[dict[str, Any]], list[str], dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    incomplete: list[str] = []
    metrics: dict[str, Any] = {"rows": 0, "columns": 0, "duplicate_rows": 0, "missing_cells": 0}
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            headers = list(reader.fieldnames or [])
            rows = list(reader)
    except (OSError, UnicodeDecodeError, csv.Error):
        return findings, ["csv_parse_failed"], metrics
    metrics.update({"rows": len(rows), "columns": len(headers)})
    if not headers or len(headers) != len(set(headers)):
        findings.append(_suite_finding(
            "DATA.SCHEMA.PROFILE", "P1", "基础数据表头为空或重复",
            location={"file": path.name, "row": 1}, actual=headers,
        ))
    fingerprints: dict[tuple[str, ...], int] = {}
    for index, row in enumerate(rows, start=2):
        values = tuple(str(row.get(header) or "").strip() for header in headers)
        if values in fingerprints:
            metrics["duplicate_rows"] += 1
            findings.append(_suite_finding(
                "DATA.SCHEMA.PROFILE", "P2", "基础数据存在完全重复行",
                location={"file": path.name, "row": index, "duplicate_of": fingerprints[values]},
            ))
        else:
            fingerprints[values] = index
        missing = [header for header, value in zip(headers, values) if not value]
        metrics["missing_cells"] += len(missing)
        if missing:
            findings.append(_suite_finding(
                "DATA.SCHEMA.PROFILE", "P2", "基础数据存在空值",
                location={"file": path.name, "row": index}, actual=missing,
            ))
        for header, value in zip(headers, values):
            lowered = header.lower()
            if lowered in {"year", "年度", "年份"} and value:
                try:
                    year = int(value)
                except ValueError:
                    year = 0
                if year < 1900 or year > 2200:
                    findings.append(_suite_finding(
                        "DATA.SCHEMA.PROFILE", "P1", "基础数据期间字段无效",
                        location={"file": path.name, "row": index, "column": header}, actual=value,
                    ))
            if lowered in {"value", "amount", "金额", "数值"} and value:
                try:
                    float(value.replace(",", ""))
                except ValueError:
                    findings.append(_suite_finding(
                        "DATA.SCHEMA.PROFILE", "P1", "基础数据数值字段无法解析",
                        location={"file": path.name, "row": index, "column": header}, actual=value,
                    ))
    if not rows:
        incomplete.append("csv_has_no_data_rows")
    return findings, incomplete, metrics


def _workbook_numeric_snapshot(path: Path) -> tuple[dict[str, float], list[str]]:
    try:
        import openpyxl

        book = openpyxl.load_workbook(path, data_only=True, read_only=True, keep_links=False)
    except Exception:  # noqa: BLE001 - external workbook may be encrypted or unsupported
        return {}, ["xlsx_value_snapshot_unavailable"]
    values: dict[str, float] = {}
    for sheet in book.worksheets:
        for row in sheet.iter_rows():
            label = str(row[0].value or "").strip().lower() if row else ""
            if not label:
                continue
            for offset, cell in enumerate(row[1:], start=2):
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    values[f"{label}|{offset}"] = float(cell.value)
    book.close()
    return values, []


def deterministic_suite_review(
    workspace_id: str,
    package_record: dict[str, Any],
    *,
    profile: str,
    standards_snapshot: dict[str, Any] | None = None,
) -> dict[str, Any]:
    package = package_record.get("payload") or {}
    findings: list[dict[str, Any]] = []
    incomplete: list[str] = package_integrity_reasons(workspace_id, package_record)
    roles = _component_roles(list(package.get("components") or []))
    missing_roles = sorted(FULL_SUITE_REQUIRED_ROLES - roles)
    if missing_roles:
        incomplete.extend(f"review_package_role_missing:{role}" for role in missing_roles)
    dimension_metrics: dict[str, dict[str, Any]] = {
        dimension: {"executed_checks": [], "finding_count": 0, "metrics": {}}
        for dimension in sorted(REVIEW_DIMENSIONS)
    }

    scope = package.get("project_scope") or {}
    if not str(scope.get("region") or "") or not str(scope.get("project_type") or ""):
        findings.append(_suite_finding(
            "COMP.REQUIREMENT.COVERAGE", "P1",
            "缺少地区或项目类型，无法冻结国家及湖北适用要求范围",
            location={"review_package_id": package_record.get("object_id")},
            actual=scope,
        ))
    dimension_metrics["compliance"]["executed_checks"].append("COMP.REQUIREMENT.COVERAGE")
    standards = standards_snapshot if isinstance(standards_snapshot, dict) else {}
    standard_issues: list[str] = []
    if not standards.get("available") or not standards.get("content_hash"):
        standard_issues.append("standards_snapshot_unavailable")
    standard_issues.extend(
        f"standard_package_incomplete:{item}"
        for item in standards.get("incomplete") or []
    )
    framework_only = {str(item) for item in standards.get("framework_only") or []}
    for standard_package in standards.get("packages") or []:
        for artifact in standard_package.get("artifacts") or []:
            if (
                str(standard_package.get("package_id") or "") in framework_only
                and not artifact.get("sha256")
            ):
                continue
            if not artifact.get("sha256"):
                standard_issues.append(
                    f"standard_artifact_hash_missing:{standard_package.get('package_id')}:{artifact.get('artifact_id')}"
                )
            if not artifact.get("publisher") or not (
                artifact.get("official_page_url") or artifact.get("source_url")
            ):
                standard_issues.append(
                    f"standard_artifact_identity_incomplete:{standard_package.get('package_id')}:{artifact.get('artifact_id')}"
                )
    incomplete.extend(standard_issues)
    dimension_metrics["compliance"]["metrics"]["standards"] = {
        "snapshot_hash": standards.get("content_hash"),
        "package_count": len(standards.get("packages") or []),
        "framework_only": sorted(framework_only),
        "issue_count": len(standard_issues),
    }

    report_texts = _package_texts(workspace_id, package, "report")
    for file_id, content in report_texts:
        placeholders = re.findall(r"待补充|待完善|TODO|TBD|XXX|\[未找到\]", content, re.I)
        if placeholders:
            findings.append(_suite_finding(
                "ARTICLE.STRUCTURE.PLACEHOLDER", "P1", "报告仍含未完成占位内容",
                location={"source_id": file_id}, actual={"count": len(placeholders)},
            ))
        seen: dict[str, int] = {}
        for index, paragraph in enumerate(re.split(r"[\r\n]+", content), start=1):
            normalized = re.sub(r"\s+", "", paragraph)
            if len(normalized) < 80:
                continue
            fingerprint = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
            if fingerprint in seen:
                findings.append(_suite_finding(
                    "ARTICLE.DUPLICATE.TEMPLATE", "P2", "报告存在长段落完全重复或模板残留",
                    location={"source_id": file_id, "paragraph": index, "duplicate_of": seen[fingerprint]},
                ))
            else:
                seen[fingerprint] = index
    dimension_metrics["article_quality"]["executed_checks"].extend([
        "ARTICLE.STRUCTURE.PLACEHOLDER", "ARTICLE.DUPLICATE.TEMPLATE",
    ])
    if not report_texts:
        incomplete.append("report_text_unavailable")

    files, analyses = _source_records(workspace_id)
    data_components = [row for row in package.get("components") or [] if row.get("role") == "base_data"]
    for component in data_components:
        if component.get("component_type") != "source_file":
            continue
        file_id = str(component.get("component_id") or "")
        analysis = analyses.get(file_id) if isinstance(analyses.get(file_id), dict) else {}
        if not analysis or not list((analysis or {}).get("locators") or []):
            findings.append(_suite_finding(
                "DATA.SCHEMA.PROFILE", "P1", "基础数据没有可审计的结构化 locator",
                location={"source_id": file_id},
            ))
        if str((analysis or {}).get("degraded_reason") or ""):
            incomplete.append(f"data_parse_degraded:{file_id}")
        record = files.get(file_id) if isinstance(files.get(file_id), dict) else {}
        path = Path(str((record or {}).get("path") or ""))
        if path.suffix.lower() == ".csv":
            csv_findings, csv_incomplete, csv_metrics = _csv_profile(path)
            for finding in csv_findings:
                finding["target_location"] = {"source_id": file_id, **(finding.get("target_location") or {})}
            findings.extend(csv_findings)
            incomplete.extend(f"data_quality:{reason}" for reason in csv_incomplete)
            dimension_metrics["data_quality"]["metrics"][file_id] = csv_metrics
    dimension_metrics["data_quality"]["executed_checks"].append("DATA.SCHEMA.PROFILE")

    evidence_components = [row for row in package.get("components") or [] if row.get("role") == "source_evidence"]
    for component in evidence_components:
        if component.get("component_type") == "source_file":
            file_id = str(component.get("component_id") or "")
            analysis = analyses.get(file_id) if isinstance(analyses.get(file_id), dict) else {}
            if not list((analysis or {}).get("locators") or []):
                findings.append(_suite_finding(
                    "SOURCE.IDENTITY.LOCATOR", "P1", "来源文件缺少可唯一解析的引用位置",
                    location={"source_id": file_id},
                ))
            if str(component.get("parse_status") or "") != "succeeded":
                incomplete.append(f"source_parse_not_succeeded:{file_id}")
    dimension_metrics["source_quality"]["executed_checks"].append("SOURCE.IDENTITY.LOCATOR")

    workbook_snapshots: dict[str, list[tuple[str, dict[str, float]]]] = {
        "finance_model": [], "finance_tables": [],
    }
    for component in package.get("components") or []:
        if component.get("role") not in {"finance_model", "finance_tables"}:
            continue
        if component.get("component_type") != "source_file":
            continue
        file_id = str(component.get("component_id") or "")
        record = files.get(file_id) if isinstance(files.get(file_id), dict) else {}
        path = Path(str((record or {}).get("path") or ""))
        suffix = path.suffix.lower()
        dimension = "financial_model" if component.get("role") == "finance_model" else "financial_tables"
        check_id = "FINMODEL.WORKBOOK.INTEGRITY" if dimension == "financial_model" else "FINTABLE.RECONCILIATION"
        if suffix not in {".xls", ".xlsx", ".xlsm"}:
            incomplete.append(f"financial_workbook_not_recalculable:{file_id}")
            continue
        if suffix == ".xlsm":
            findings.append(_suite_finding(
                check_id, "P1", "XLSM 含宏容器；本系统只做静态识别且绝不执行宏",
                location={"source_id": file_id}, remediation="提供无宏可重算副本或由专业人员独立核验宏逻辑",
            ))
        scanned, missing, _metrics = rules.scan_xlsx(path, deep=profile == "deep")
        for finding in scanned:
            finding["rule_id"] = check_id
            finding["category"] = dimension
            finding["review_area"] = dimension
            finding["target_location"] = {"source_id": file_id, **(finding.get("target_location") or {})}
        findings.extend(scanned)
        incomplete.extend(f"{dimension}:{reason}" for reason in missing)
        snapshot, snapshot_missing = _workbook_numeric_snapshot(path)
        workbook_snapshots[str(component.get("role") or "")].append((file_id, snapshot))
        incomplete.extend(f"{dimension}:{reason}" for reason in snapshot_missing)
        dimension_metrics[dimension]["metrics"][file_id] = {
            **_metrics,
            "numeric_snapshot_cells": len(snapshot),
        }
        if profile == "deep" and suffix in {".xlsx", ".xlsm"}:
            recalculated, recalc_missing, _recalc_metrics = rules.recalculate_xlsx(path)
            for finding in recalculated:
                finding["rule_id"] = check_id
                finding["category"] = dimension
                finding["review_area"] = dimension
                finding["target_location"] = {"source_id": file_id, **(finding.get("target_location") or {})}
            findings.extend(recalculated)
            incomplete.extend(f"{dimension}:{reason}" for reason in recalc_missing)
        dimension_metrics[dimension]["executed_checks"].append(check_id)

    for model_id, model_values in workbook_snapshots["finance_model"]:
        for table_id, table_values in workbook_snapshots["finance_tables"]:
            shared = sorted(set(model_values).intersection(table_values))
            mismatches = [
                key for key in shared
                if abs(model_values[key] - table_values[key]) > max(0.01, abs(model_values[key]) * 1e-9)
            ]
            if mismatches:
                findings.append(_suite_finding(
                    "FINTABLE.RECONCILIATION", "P1", "财务模型与财务表同名期间数值不一致",
                    location={"finance_model_source_id": model_id, "finance_tables_source_id": table_id},
                    actual={
                        key: {"model": model_values[key], "tables": table_values[key]}
                        for key in mismatches[:50]
                    },
                ))
            dimension_metrics["financial_tables"]["metrics"][f"{model_id}:{table_id}"] = {
                "shared_numeric_cells": len(shared),
                "mismatch_count": len(mismatches),
            }

    required_groups = {
        "市场需求": ("市场", "需求"),
        "建设方案": ("建设方案", "技术方案", "工程方案"),
        "投资融资": ("投资", "融资", "资金筹措"),
        "财务评价": ("财务", "内部收益率", "现金流"),
        "风险结论": ("风险", "结论", "建议"),
    }
    combined_text = "\n".join(text for _file_id, text in report_texts)
    missing_groups = [
        name for name, terms in required_groups.items()
        if not any(term in combined_text for term in terms)
    ]
    if missing_groups:
        findings.append(_suite_finding(
            "FEASIBILITY.STRUCTURE.COVERAGE", "P1", "可行性研究论证链章节覆盖不足",
            location={"review_package_id": package_record.get("object_id")}, actual=missing_groups,
        ))
    dimension_metrics["feasibility"]["executed_checks"].append("FEASIBILITY.STRUCTURE.COVERAGE")

    for dimension in dimension_metrics:
        dimension_metrics[dimension]["finding_count"] = sum(
            1 for row in findings if row.get("review_area") == dimension
        )
    unique = {str(row.get("finding_id") or sha256_json(row)): row for row in findings}
    ordered = sorted(
        unique.values(),
        key=lambda row: (SEVERITY_ORDER.get(str(row.get("severity") or ""), 9), str(row.get("finding_id") or "")),
    )
    return {
        "findings": ordered,
        "incomplete_reasons": sorted(set(incomplete)),
        "dimension_metrics": dimension_metrics,
        "full_suite": not missing_roles,
    }


def _assessment_evidence(
    workspace_id: str,
    package: dict[str, Any],
    rows: Any,
) -> tuple[list[dict[str, Any]], dict[str, Any] | None]:
    from lvke_mcp.adapters.source_files_repository import SourceFileError, resolve_citation_fragment

    components = {
        str(row.get("component_id") or ""): row
        for row in (package.get("payload") or {}).get("components") or []
    }
    verified: list[dict[str, Any]] = []
    for raw in rows or []:
        source_id = str(raw.get("source_id") or "")
        if source_id not in components:
            return [], _blocked("assessment_evidence_not_in_package", "Assessment 证据不属于冻结套件")
        component = components[source_id]
        if component.get("component_type") != "source_file":
            target_spec = component.get("target_spec")
            if not isinstance(target_spec, dict):
                return [], _blocked("assessment_internal_evidence_invalid", "内部证据缺少冻结 target_spec")
            internal, reasons = _internal_component(workspace_id, target_spec)
            if reasons or internal is None:
                return [], _blocked(
                    reasons[0] if reasons else "assessment_internal_evidence_unavailable",
                    "内部证据对象无法重新解析",
                )
            if str(internal.get("content_hash") or "") != str(component.get("content_hash") or ""):
                return [], _blocked("assessment_internal_evidence_hash_mismatch", "内部证据对象已变化")
            supplied_hash = _hash(str(raw.get("source_hash") or ""))
            if supplied_hash != str(component.get("content_hash") or ""):
                return [], _blocked("assessment_internal_evidence_hash_mismatch", "内部证据 source_hash 不匹配")
            from .target_resolve import _resolve_target
            from lvke_mcp.servers.lvke_deliverable_review.contracts import normalize_target

            resolved, blockers = _resolve_target(workspace_id, normalize_target(target_spec))
            if blockers or resolved is None:
                return [], _blocked(blockers[0], "内部证据对象无法解析")
            locator = raw.get("locator")
            pointer = (
                str(locator.get("json_pointer") or "")
                if isinstance(locator, dict)
                else str(locator or "").removeprefix("json:")
            )
            if not pointer.startswith("/"):
                return [], _blocked("assessment_internal_locator_invalid", "内部对象 locator 必须是 JSON Pointer")
            value: Any = resolved.get("snapshot")
            try:
                for token in pointer.split("/")[1:]:
                    token = token.replace("~1", "/").replace("~0", "~")
                    value = value[int(token)] if isinstance(value, list) else value[token]
            except (KeyError, IndexError, TypeError, ValueError):
                return [], _blocked("assessment_internal_locator_not_found", "内部对象 JSON Pointer 无法唯一解析")
            fragment = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
            fragment_hash = "sha256:" + hashlib.sha256(fragment.encode("utf-8")).hexdigest()
            supplied_fragment = str(raw.get("fragment_text") or "")
            supplied_fragment_hash = _hash(str(raw.get("fragment_hash") or ""))
            if supplied_fragment and supplied_fragment != fragment:
                return [], _blocked("assessment_fragment_mismatch", "内部对象 fragment_text 与解析结果不一致")
            if supplied_fragment_hash and supplied_fragment_hash != fragment_hash:
                return [], _blocked("assessment_fragment_hash_mismatch", "内部对象 fragment_hash 与解析结果不一致")
            verified.append({
                "source_id": source_id,
                "source_hash": supplied_hash,
                "locator": {"json_pointer": pointer},
                "fragment_text": fragment,
                "fragment_hash": fragment_hash,
                "source_kind": "immutable_object",
            })
            continue
        try:
            result = resolve_citation_fragment(
                workspace_id,
                source_id=source_id,
                locator=raw.get("locator"),
                source_hash=raw.get("source_hash"),
                supplied_fragment=raw.get("fragment_text", ""),
                supplied_fragment_hash=raw.get("fragment_hash", ""),
            )
        except SourceFileError as exc:
            return [], _blocked(str(exc.detail.get("code") or "assessment_evidence_invalid"), str(exc.detail.get("message") or "Assessment 证据无效"))
        verified.append({**result, "source_id": source_id})
    return verified, None


def submit_assessment(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        review_id = str(args.get("review_id") or "")
        package_id = str(args.get("review_package_id") or "")
        package = _verified_record(REVIEW_PACKAGE_STORE, workspace_id, package_id)
        if package is None:
            return _blocked("review_package_not_found", "审查套件不存在或完整性无效")
        dimension = str(args.get("dimension") or "")
        if dimension not in REVIEW_DIMENSIONS:
            return _blocked("review_dimension_invalid", "审查维度不受支持")
        status = str(args.get("status") or "")
        if status not in DIMENSION_STATUSES:
            return _blocked("review_dimension_status_invalid", "审查维度状态不受支持")
        try:
            from .events import _project

            state = _project(workspace_id, review_id)
        except ValueError:
            return _blocked("review_not_found", _message("review_not_found"))
        if str((state.get("target") or {}).get("target_type") or "") != "review_package" or str((state.get("target") or {}).get("target_id") or "") != package_id:
            return _blocked("assessment_review_package_mismatch", "Assessment 与 Review 绑定的套件不一致")
        findings: list[dict[str, Any]] = []
        for raw in args.get("findings") or []:
            check_id = str(raw.get("check_id") or "")
            spec = CHECK_CATALOG.get(check_id)
            if spec is None or spec["dimension"] != dimension or spec["kind"] != "semantic":
                return _blocked("assessment_check_id_invalid", "Assessment finding 必须使用当前维度已登记的 semantic check_id")
            evidence, error = _assessment_evidence(workspace_id, package, raw.get("evidence") or [])
            if error is not None:
                return error
            missing_reason = str(raw.get("missing_evidence_reason") or "").strip()
            if not evidence and not missing_reason:
                return _blocked("assessment_evidence_required", "每条语义 finding 必须绑定可验证证据或明确缺证原因")
            severity = str(raw.get("severity") or "")
            if severity not in {"P0", "P1", "P2", "P3"}:
                return _blocked("assessment_severity_invalid", "Assessment finding 严重度无效")
            finding = _suite_finding(
                check_id,
                severity,
                str(raw.get("message") or ""),
                location=deepcopy(raw.get("target_location") or {}),
                actual=deepcopy(raw.get("actual")),
                remediation=str(raw.get("remediation") or "补充材料或修订后由同一领域独立复审"),
            )
            finding.update({
                "evidence": evidence,
                "missing_evidence_reason": missing_reason,
                "manual_review_required": True,
                "assessment_source": "independent_agent",
            })
            findings.append(finding)
        if status == "passed" and findings:
            return _blocked("assessment_pass_with_findings", "passed Assessment 不能同时提交 findings")
        if status == "failed" and not findings:
            return _blocked("assessment_failed_without_findings", "failed Assessment 必须至少提交一个 finding")
        package_profile = str((package.get("payload") or {}).get("review_profile") or "standard")
        required_check_ids = {
            check_id
            for check_id, spec in CHECK_CATALOG.items()
            if spec["dimension"] == dimension
            and spec["kind"] == "semantic"
            and (check_id != "ARTICLE.VISUAL.LAYOUT" or package_profile == "deep")
        }
        coverage = deepcopy(args.get("coverage") or {})
        covered_check_ids = {
            str(item) for item in coverage.get("checked_check_ids") or [] if str(item)
        }
        if status in {"passed", "failed"} and not required_check_ids.issubset(covered_check_ids):
            return _blocked(
                "assessment_coverage_incomplete",
                "Assessment 必须显式登记当前 profile 的全部 semantic check_id",
                missing_check_ids=sorted(required_check_ids - covered_check_ids),
            )
        reviewer_context_id = str(args.get("reviewer_context_id") or "").strip()
        if not reviewer_context_id:
            return _blocked("assessment_reviewer_context_id_required", "Assessment 必须记录独立 reviewer_context_id")
        for existing in SUITE_ASSESSMENT_STORE.list(workspace_id):
            existing_payload = existing.get("payload") or {}
            existing_context = (existing_payload.get("reviewer_context") or {}).get("reviewer_context_id")
            if (
                str(existing_payload.get("review_id") or "") == review_id
                and str(existing_payload.get("dimension") or "") != dimension
                and str(existing_context or "") == reviewer_context_id
            ):
                return _blocked(
                    "assessment_context_not_independent",
                    "同一 reviewer_context_id 不得承担多个七域审查通道",
                )
        assessment_payload = {
            "review_id": review_id,
            "review_package_id": package_id,
            "review_package_hash": package["content_hash"],
            "dimension": dimension,
            "status": status,
            "coverage": coverage,
            "findings": findings,
            "limitations": list(args.get("limitations") or []),
            "reviewer_context": {
                "skill": str(args.get("skill") or ""),
                "skill_version": str(args.get("skill_version") or ""),
                "model": str(args.get("model") or ""),
                "model_version": str(args.get("model_version") or ""),
                "execution_environment": str(args.get("execution_environment") or "controlled_current_environment"),
                "independent_context": bool(args.get("independent_context")),
                "reviewer_context_id": reviewer_context_id,
            },
        }
        if not assessment_payload["reviewer_context"]["independent_context"]:
            return _blocked("independent_review_context_required", "正式七域 Assessment 必须来自独立审查上下文")
        if not all(assessment_payload["reviewer_context"].get(key) for key in ("skill", "skill_version", "model", "model_version")):
            return _blocked("assessment_reviewer_metadata_required", "Assessment 必须记录 Skill、模型及版本")
        record = SUITE_ASSESSMENT_STORE.put(
            workspace_id,
            assessment_payload,
            producer="lvke-deliverable-review.review_submit_assessment",
            source_ids=[review_id, package_id],
            basis={
                "review_id": review_id,
                "review_package_hash": package["content_hash"],
                "dimension": dimension,
                "reviewer_context": assessment_payload["reviewer_context"],
            },
            schema_version="ReviewAssessment.v1",
        )
        STORE.append(workspace_id, review_id, "suite_assessment_submitted", {
            "assessment_id": record["object_id"],
            "assessment_hash": record["content_hash"],
            "dimension": dimension,
            "status": status,
            "findings": findings,
        })
        return _ok(
            review_id=review_id,
            review_assessment_id=record["object_id"],
            dimension=dimension,
            dimension_status=status,
            finding_count=len(findings),
            resource_uris=[record["resource_uri"]],
            warnings=["该结果是结构化 Agent 专业审查，不等同于法律或执业签署"],
            blockers=[],
            next_actions=["领域审查完成后调用 review_confirm_dimension"],
        )
    return _write("review_submit_assessment", args, execute)


def confirm_dimension(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        review_id = str(args.get("review_id") or "")
        dimension = str(args.get("dimension") or "")
        if dimension not in REVIEW_DIMENSIONS:
            return _blocked("review_dimension_invalid", "审查维度不受支持")
        assessments = [
            row for row in SUITE_ASSESSMENT_STORE.list(workspace_id)
            if str((row.get("payload") or {}).get("review_id") or "") == review_id
            and str((row.get("payload") or {}).get("dimension") or "") == dimension
        ]
        if not assessments:
            return _blocked("review_dimension_assessment_required", "该维度尚无独立 Assessment")
        assessment = sorted(assessments, key=lambda row: str(row.get("created_at") or ""))[-1]
        payload = {
            "review_id": review_id,
            "dimension": dimension,
            "assessment_id": assessment["object_id"],
            "assessment_hash": assessment["content_hash"],
            "role_declaration": str(args.get("role_declaration") or ""),
            "review_statement": str(args.get("review_statement") or ""),
            "limitations_accepted": list(args.get("limitations_accepted") or []),
            "identity_or_credential_verified": False,
        }
        if not payload["role_declaration"] or not payload["review_statement"]:
            return _blocked("dimension_confirmation_statement_required", "领域确认必须提供角色声明和审查意见")
        record = DIMENSION_CONFIRMATION_STORE.put(
            workspace_id,
            payload,
            producer="lvke-deliverable-review.review_confirm_dimension",
            source_ids=[review_id, assessment["object_id"]],
            basis=payload,
            schema_version="ReviewDimensionConfirmation.v1",
        )
        STORE.append(workspace_id, review_id, "suite_dimension_confirmed", {
            **payload,
            "dimension_confirmation_id": record["object_id"],
            "confirmation_hash": record["content_hash"],
        })
        return _ok(
            review_id=review_id,
            dimension=dimension,
            dimension_confirmation_id=record["object_id"],
            identity_or_credential_verified=False,
            resource_uris=[record["resource_uri"]],
            warnings=["角色确认是审查责任声明，不是身份、资质或电子签名认证"],
            blockers=[],
            next_actions=["全部必审维度确认后调用 review_finalize"],
        )
    return _write("review_confirm_dimension", args, execute)


def _latest_records(store: Any, workspace_id: str, review_id: str, key: str) -> dict[str, dict[str, Any]]:
    selected: dict[str, dict[str, Any]] = {}
    for record in sorted(store.list(workspace_id), key=lambda row: str(row.get("created_at") or "")):
        payload = record.get("payload") or {}
        if str(payload.get("review_id") or "") == review_id:
            selected[str(payload.get(key) or "")] = record
    return selected


def _dimension_results(
    state: dict[str, Any],
    assessments: dict[str, dict[str, Any]],
    confirmations: dict[str, dict[str, Any]],
    *,
    require_semantic: bool,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    incomplete_reasons = [str(item) for item in state.get("incomplete_reasons") or []]
    dimension_prefixes = {
        "compliance": ("standard_", "standards_", "COMP.", "review_package_role_missing"),
        "article_quality": ("report_", "article_"),
        "data_quality": ("data_", "csv_"),
        "source_quality": ("source_", "citation_", "extraction_"),
        "financial_model": ("financial_model:", "xlsx_", "libreoffice_"),
        "financial_tables": ("financial_tables:",),
        "feasibility": ("feasibility_",),
    }
    for dimension in sorted(REVIEW_DIMENSIONS):
        projected_findings = [
            row for row in state.get("findings") or []
            if row.get("review_area") == dimension
        ]
        active_findings = [
            row for row in projected_findings
            if row.get("status") not in {"resolved", "rejected", "superseded", "waived"}
        ]
        deterministic_findings = [
            row for row in active_findings
            if row.get("manual_review_required") is not True
        ]
        semantic_findings = [
            row for row in active_findings
            if row.get("manual_review_required") is True
        ]
        assessment_record = assessments.get(dimension)
        assessment = (assessment_record or {}).get("payload") or {}
        confirmation_record = confirmations.get(dimension)
        dimension_incomplete_reasons = [
            reason
            for reason in incomplete_reasons
            if reason.startswith(dimension_prefixes[dimension])
        ]
        deterministic_incomplete = bool(dimension_incomplete_reasons)
        deterministic_status = "failed" if any(finding_blocks(row) for row in deterministic_findings) else "passed"
        if deterministic_status == "passed" and deterministic_incomplete:
            deterministic_status = "incomplete"
        assessment_present = bool(assessment_record)
        semantic_status = str(assessment.get("status") or ("incomplete" if require_semantic else "not_applicable"))
        if semantic_status == "failed" and not semantic_findings and assessment.get("findings"):
            semantic_status = "passed"
        # Quick is a deterministic preview only.  Without an Assessment it
        # must remain incomplete and unconfirmed instead of looking passed.
        confirmed = bool(confirmation_record) if require_semantic else assessment_present
        if deterministic_status == "failed" or semantic_status == "failed":
            status = "failed"
        elif (
            deterministic_status == "incomplete"
            or semantic_status in {"incomplete", "not_determinable"}
            or not confirmed
            or (not require_semantic and not assessment_present)
        ):
            status = "incomplete" if semantic_status != "not_determinable" else "not_determinable"
        else:
            status = "passed"
        compliance_status = "professional_determination_required"
        if dimension == "compliance":
            if status == "failed":
                compliance_status = "nonconformity_found"
            elif status in {"incomplete", "not_determinable"}:
                compliance_status = "evidence_incomplete"
            elif status == "passed":
                compliance_status = "conforms_to_checked_requirements"
        results.append({
            "dimension": dimension,
            "status": status,
            "deterministic_status": deterministic_status,
            "semantic_status": semantic_status,
            "assessment_id": str((assessment_record or {}).get("object_id") or ""),
            "confirmation_id": str((confirmation_record or {}).get("object_id") or ""),
            # 责任声明与审查意见随结果一并给出：``review_confirm_dimension`` 强制
            # 要求这两项非空，只回一个 confirmation_id 会让消费方为了显示"谁按什么
            # 责任确认的"再逐条回查确认对象。``limitations_accepted`` 同理——它是
            # 判断"限制项是否已被接受"的唯一依据。
            "role_declaration": str(
                ((confirmation_record or {}).get("payload") or {}).get("role_declaration") or ""
            ),
            "review_statement": str(
                ((confirmation_record or {}).get("payload") or {}).get("review_statement") or ""
            ),
            "limitations_accepted": list(
                ((confirmation_record or {}).get("payload") or {}).get("limitations_accepted") or []
            ),
            "confirmed_at": str((confirmation_record or {}).get("created_at") or ""),
            "role_confirmed": confirmed,
            "identity_or_credential_verified": False,
            "compliance_status": compliance_status,
            # 逐维度列出使其 incomplete 的原因。只给 status=incomplete 会让消费方
            # 无法区分"审查没跑完"与"材料结构性缺项"，只能靠猜——粗粒度状态掩盖
            # 根因正是排查方向跑偏的来源。
            "incomplete_reasons": dimension_incomplete_reasons,
            "finding_count": len(active_findings),
            "limitations": [
                *list(assessment.get("limitations") or []),
                *(["quick_preview_without_semantic_assessment"] if not require_semantic and not assessment_present else []),
            ],
        })
    return results


def get_dimension(args: dict[str, Any]) -> dict[str, Any]:
    workspace_id = str(args.get("workspace_id") or "")
    review_id = str(args.get("review_id") or "")
    dimension = str(args.get("dimension") or "")
    try:
        require_safe_id(workspace_id, "workspace_id")
        require_safe_id(review_id, "review_id")
    except ValueError as exc:
        return _blocked(str(exc), _message(str(exc)))
    if dimension not in REVIEW_DIMENSIONS:
        return _blocked("review_dimension_invalid", "审查维度不受支持")
    try:
        from .events import _project

        state = _project(workspace_id, review_id)
    except ValueError:
        return _blocked("review_not_found", _message("review_not_found"))
    assessments = _latest_records(SUITE_ASSESSMENT_STORE, workspace_id, review_id, "dimension")
    confirmations = _latest_records(DIMENSION_CONFIRMATION_STORE, workspace_id, review_id, "dimension")
    profile = str(state.get("review_profile") or state.get("mode") or "quick")
    result = next(row for row in _dimension_results(
        state, assessments, confirmations, require_semantic=profile in {"standard", "deep"},
    ) if row["dimension"] == dimension)
    findings = [row for row in state.get("findings") or [] if row.get("review_area") == dimension]
    assessment = (assessments.get(dimension) or {}).get("payload") or {}
    return _ok(
        review_id=review_id,
        dimension_result=result,
        findings=[*findings, *list(assessment.get("findings") or [])],
        coverage=deepcopy(assessment.get("coverage") or {}),
        resource_uris=[str((assessments.get(dimension) or {}).get("resource_uri") or "")],
        warnings=[], blockers=[], next_actions=[],
    )


def finalize(args: dict[str, Any]) -> dict[str, Any]:
    def execute(workspace_id: str) -> dict[str, Any]:
        review_id = str(args.get("review_id") or "")
        try:
            from .events import _project

            state = _project(workspace_id, review_id)
        except ValueError:
            return _blocked("review_not_found", _message("review_not_found"))
        if str((state.get("target") or {}).get("target_type") or "") != "review_package":
            return _blocked("review_package_target_required", "七域 finalize 仅适用于 ReviewPackage")
        package_id = str((state.get("target") or {}).get("target_id") or "")
        package = _verified_record(REVIEW_PACKAGE_STORE, workspace_id, package_id)
        if package is None:
            return _blocked("review_package_not_found", "审查套件不存在或完整性无效")
        integrity_reasons = package_integrity_reasons(workspace_id, package)
        if integrity_reasons:
            return _blocked(
                integrity_reasons[0],
                "finalize 前 ReviewPackage 完整性或正式谱系校验失败",
                blockers=integrity_reasons,
            )
        profile = str(state.get("review_profile") or state.get("mode") or "quick")
        require_semantic = profile in {"standard", "deep"}
        assessments = _latest_records(SUITE_ASSESSMENT_STORE, workspace_id, review_id, "dimension")
        confirmations = _latest_records(DIMENSION_CONFIRMATION_STORE, workspace_id, review_id, "dimension")
        dimensions = _dimension_results(
            state, assessments, confirmations, require_semantic=require_semantic,
        )
        blockers: list[str] = []
        package_payload = package.get("payload") or {}
        if not package_payload.get("full_suite"):
            blockers.extend(
                f"review_package_role_missing:{role}"
                for role in package_payload.get("missing_required_roles") or []
            )
        if profile == "quick":
            blockers.append("quick_profile_not_formal_suite_review")
        blockers.extend(f"review_incomplete:{item}" for item in state.get("incomplete_reasons") or [])
        required_extractions = {
            str(item) for item in package_payload.get("extraction_confirmation_required") or []
        }
        confirmed_extractions = {
            str(item.get("source_id") or "")
            for record in EXTRACTION_CONFIRMATION_STORE.list(workspace_id)
            if str((record.get("payload") or {}).get("review_package_id") or "") == package_id
            and str((record.get("payload") or {}).get("review_package_hash") or "") == package.get("content_hash")
            for item in ((record.get("payload") or {}).get("confirmations") or [])
        }
        blockers.extend(
            f"extraction_confirmation_missing:{source_id}"
            for source_id in sorted(required_extractions - confirmed_extractions)
        )
        for result in dimensions:
            if result["status"] in {"failed", "incomplete", "not_determinable"}:
                blockers.append(f"review_dimension_{result['status']}:{result['dimension']}")
            if require_semantic and not result["role_confirmed"]:
                blockers.append(f"review_dimension_unconfirmed:{result['dimension']}")
        active = [
            row for row in state.get("findings") or []
            if row.get("status") not in {"resolved", "rejected", "superseded", "waived"}
        ]
        blockers.extend(f"blocking_finding:{row.get('finding_id')}" for row in active if finding_blocks(row))
        blockers = sorted(set(blockers))
        overall = "pass" if not blockers else (
            "incomplete" if any("incomplete" in item or "not_determinable" in item or "missing" in item or "unconfirmed" in item for item in blockers)
            else "fail"
        )
        dimension_records: list[dict[str, Any]] = []
        for result in dimensions:
            dimension = str(result.get("dimension") or "")
            assessment_record = assessments.get(dimension) or {}
            confirmation_record = confirmations.get(dimension) or {}
            dimension_payload = {
                "schema_version": "ReviewDimensionResult.v1",
                "review_id": review_id,
                "review_package_id": package_id,
                "review_package_hash": package["content_hash"],
                "review_profile": profile,
                **deepcopy(result),
                "deterministic_review_event_chain_hash": str(state.get("event_chain_hash") or ""),
            }
            dimension_records.append(DIMENSION_RESULT_STORE.put(
                workspace_id,
                dimension_payload,
                producer="lvke-deliverable-review.review_finalize",
                source_ids=[
                    review_id,
                    package_id,
                    *(
                        [str(assessment_record.get("object_id") or "")]
                        if assessment_record.get("object_id")
                        else []
                    ),
                    *(
                        [str(confirmation_record.get("object_id") or "")]
                        if confirmation_record.get("object_id")
                        else []
                    ),
                ],
                basis=dimension_payload,
                schema_version="ReviewDimensionResult.v1",
            ))
        dossier_payload = {
            "schema_version": "ReviewDossier.v2",
            "review_id": review_id,
            "review_package_id": package_id,
            "review_package_hash": package["content_hash"],
            "review_mode": str(package_payload.get("review_mode") or ""),
            "review_profile": profile,
            "full_suite": bool(package_payload.get("full_suite")),
            "formal_suite_review_complete": overall == "pass" and profile in {"standard", "deep"},
            "dimension_results": dimensions,
            "dimension_result_ids": sorted(
                str(row.get("object_id") or "") for row in dimension_records
            ),
            "assessment_ids": sorted(str(row.get("object_id") or "") for row in assessments.values()),
            "dimension_confirmation_ids": sorted(str(row.get("object_id") or "") for row in confirmations.values()),
            "deterministic_review_event_chain_hash": str(state.get("event_chain_hash") or ""),
            "overall_verdict": overall,
            "hard_gate_blockers": blockers,
            "standards_snapshot": deepcopy(state.get("standards") or {}),
            "standards_snapshot_hash": str((state.get("standards") or {}).get("content_hash") or ""),
            "deterministic_coverage": deepcopy(state.get("coverage") or {}),
            "professional_signoff_determined": False,
        }
        record = DOSSIER_STORE.put(
            workspace_id,
            dossier_payload,
            producer="lvke-deliverable-review.review_finalize",
            source_ids=[
                review_id,
                package_id,
                *dossier_payload["assessment_ids"],
                *dossier_payload["dimension_confirmation_ids"],
                *dossier_payload["dimension_result_ids"],
            ],
            basis={key: dossier_payload[key] for key in (
                "review_id", "review_package_hash", "dimension_results",
                "dimension_result_ids", "deterministic_review_event_chain_hash",
                "overall_verdict",
            )},
            schema_version="ReviewDossier.v2",
        )
        STORE.append(workspace_id, review_id, "suite_finalized", {
            "dossier_id": record["object_id"],
            "dossier_hash": record["content_hash"],
            "dimension_results": dimensions,
            "dimension_result_ids": dossier_payload["dimension_result_ids"],
            "overall_verdict": overall,
            "hard_gate_blockers": blockers,
            "formal_suite_review_complete": dossier_payload["formal_suite_review_complete"],
        })
        retest_result = None
        if dossier_payload["formal_suite_review_complete"]:
            from .suite_retest import complete_pending_suite_retest

            retest_result = complete_pending_suite_retest(
                workspace_id,
                review_id,
                check_catalog=CHECK_CATALOG,
            )
        return _ok(
            status="ok" if overall == "pass" else "incomplete" if overall == "incomplete" else "partial",
            review_id=review_id,
            review_dossier_id=record["object_id"],
            dimension_results=dimensions,
            dimension_result_ids=dossier_payload["dimension_result_ids"],
            overall_verdict=overall,
            formal_suite_review_complete=dossier_payload["formal_suite_review_complete"],
            retest_result=retest_result,
            professional_signoff_determined=False,
            resource_uris=[
                record["resource_uri"],
                *(str(row.get("resource_uri") or "") for row in dimension_records),
            ],
            warnings=["结论不代表法律批准、执业签章、身份认证或电子签名"],
            blockers=blockers,
            next_actions=[] if overall == "pass" else ["处理硬门禁问题并创建新套件/Retest"],
        )
    return _write("review_finalize", args, execute)


__all__ = [
    "Any",
    "CHECK_CATALOG",
    "COMPLIANCE_STATUSES",
    "DIMENSION_CONFIRMATION_STORE",
    "DIMENSION_RESULT_STORE",
    "DIMENSION_STATUSES",
    "DOSSIER_STORE",
    "EXTRACTION_CONFIRMATION_STORE",
    "FULL_SUITE_REQUIRED_ROLES",
    "FormalLineageError",
    "PACKAGE_DRAFT_STORE",
    "Path",
    "REVIEW_COMPONENT_ROLES",
    "REVIEW_DIMENSIONS",
    "REVIEW_MODES",
    "REVIEW_PACKAGE_STORE",
    "REVIEW_PROFILES",
    "SEVERITY_ORDER",
    "STORE",
    "SUITE_ASSESSMENT_STORE",
    "_ROLE_TERMS",
    "_assessment_evidence",
    "_blocked",
    "_build_internal_component",
    "_component_roles",
    "_csv_profile",
    "_dimension_results",
    "_hash",
    "_internal_component",
    "_internal_package_lineage",
    "_latest_records",
    "_message",
    "_ok",
    "_package_texts",
    "_source_records",
    "_suggest_role",
    "_suite_finding",
    "_validate_internal_package_lineage",
    "_validate_package_integrity",
    "_verified_record",
    "_verified_source_component",
    "_workbook_numeric_snapshot",
    "_write",
    "confirm_dimension",
    "confirm_extraction",
    "confirm_package",
    "csv",
    "deepcopy",
    "deterministic_suite_review",
    "finalize",
    "finding_blocks",
    "get_dimension",
    "get_package",
    "hashlib",
    "json",
    "package_integrity_reasons",
    "prepare_package",
    "re",
    "require_safe_id",
    "rules",
    "sha256_json",
    "submit_assessment",
    "utc_now",
]
